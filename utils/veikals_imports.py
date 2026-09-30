"""
Imports no esošā Jelgavas veikala inventarizācijas Excel faila
(piem., "inventarizācija Augusts 2026.xlsx", lapa "01112025").

Kolonnas (galviņa 3. rindā):
  A Product · B thickness · C width · D lenght · E pcs in pack · F packs
  I Pašizmaksa (€/m³, gabalprecēm €/gab) · L pack number · M Piegādātājs
  N pēdējā inventerizācija (gab) · Q Pārdošanas PPR (pavadzīme)
  U Pārdošanas cena m3 (gabalprecēm €/gab) · W Pārdošanas cena m2
  Z aktuālākā inventerizācija (gab) · AH KOMENTARI 2

Atlikums = Z (ja aizpildīts), citādi N.
Lapa "Granulas": A produkts · B maisi · E cena paš/maiss · H pvz · I pēd. inv. · L pārd. cena · Q akt. inv.
"""
import openpyxl

KOL = {"produkts": 0, "biezums": 1, "platums": 2, "garums": 3, "gab_paka": 4, "pakas": 5,
       "pasizmaksa": 8, "pakas_nr": 11, "piegadatajs": 12, "ped_inv": 13, "ppr": 16,
       "cena_m3": 20, "cena_m2": 22, "akt_inv": 25, "komentars": 33}


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", ".").strip())
        except ValueError:
            return None
    return None


def _txt(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def _atrast_galvinu(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        vals = [str(v).strip().lower() for v in row if v is not None]
        if "thickness" in vals and "product" in vals:
            return i
    return None


def lapas_saraksts(fails):
    wb = openpyxl.load_workbook(fails, read_only=True, data_only=True)
    res = []
    for ws in wb.worksheets:
        if ws.title.lower() == "granulas":
            res.append((ws.title, "granulas"))
        elif _atrast_galvinu(ws):
            res.append((ws.title, "pakas"))
    wb.close()
    return res


def nolasit_pakas(fails, lapa):
    """Atgriež sarakstu ar pakas dict (derīgs `add_paka`) un izlaisto rindu skaitu."""
    wb = openpyxl.load_workbook(fails, data_only=True)
    ws = wb[lapa]
    galvina = _atrast_galvinu(ws)
    if galvina is None:
        raise ValueError(f"Lapā '{lapa}' nav atrasta galviņa (Product / thickness)")
    pakas, izlaistas = [], 0
    for row in ws.iter_rows(min_row=galvina + 1, values_only=True):
        row = list(row) + [None] * (40 - len(row))
        produkts = _txt(row[KOL["produkts"]])
        b, p, g = (_num(row[KOL[k]]) for k in ("biezums", "platums", "garums"))
        akt, ped = _num(row[KOL["akt_inv"]]), _num(row[KOL["ped_inv"]])
        atlikums = akt if akt is not None else ped
        if atlikums is None:
            atlikums = _num(row[KOL["gab_paka"]]) or 0
        ir_izmeri = bool(b and g)  # mietiem platuma nav (b = diametrs)
        if not produkts and not ir_izmeri:
            # Tukšas vai kopsavilkuma rindas faila beigās
            if any(v is not None for v in row[:14]):
                izlaistas += 1
            continue
        gab_paka = _num(row[KOL["gab_paka"]])
        pakas_sk = _num(row[KOL["pakas"]]) or 1
        gab_sakuma = int((gab_paka or atlikums) * pakas_sk)
        pasizm = _num(row[KOL["pasizmaksa"]])
        cena_u = _num(row[KOL["cena_m3"]])
        paka = {
            "produkts": produkts or "Nezināms",
            "biezums": b if ir_izmeri else None,
            "platums": p if (ir_izmeri and p) else None,
            "garums": g if ir_izmeri else None,
            "gab_sakuma": gab_sakuma,
            "gab_atlikums": int(atlikums),
            "pasizmaksa_m3": pasizm if ir_izmeri else None,
            "pasizmaksa_gab": None if ir_izmeri else pasizm,
            "cena_m3": round(cena_u, 2) if (ir_izmeri and cena_u) else None,
            "cena_m2": _num(row[KOL["cena_m2"]]) if ir_izmeri else None,
            "cena_gab": None if ir_izmeri else cena_u,
            "pakas_nr": _txt(row[KOL["pakas_nr"]]),
            "piegadatajs": _txt(row[KOL["piegadatajs"]]),
            "pavadzime": _txt(row[KOL["ppr"]]),
            "piezimes": _txt(row[KOL["komentars"]]),
        }
        pakas.append(paka)
    wb.close()
    return pakas, izlaistas


def nolasit_granulas(fails, lapa="Granulas"):
    wb = openpyxl.load_workbook(fails, data_only=True)
    ws = wb[lapa]
    pakas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * (25 - len(row))
        produkts = _txt(row[0])
        if not produkts:
            continue
        akt, ped = _num(row[16]), _num(row[8])
        atlikums = akt if akt is not None else (ped if ped is not None else (_num(row[1]) or 0))
        pakas.append({
            "produkts": produkts,
            "gab_sakuma": int(_num(row[1]) or atlikums),
            "gab_atlikums": int(atlikums),
            "pasizmaksa_gab": _num(row[4]),
            "cena_gab": _num(row[11]),
            "pavadzime": _txt(row[7]),
            "piegadatajs": "DVK solutions",
            "piezimes": "maisi (1 paletē)",
        })
    wb.close()
    return pakas
