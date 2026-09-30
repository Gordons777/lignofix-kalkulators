"""
Jelgavas veikala Excel eksporti (openpyxl, ar formulām):
  • pavadzime_xlsx      — pārdošanas pavadzīme / čeks vienai pārdošanai
  • atlikumi_xlsx       — noliktavas atlikumi (inventarizācijas lapa)
  • atskaite_xlsx       — pārdošanas atskaite periodā + kopsavilkums pa piegādātājiem
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.veikals_calc import REVERSA_TEKSTS

ZILS = PatternFill("solid", fgColor="1F4E78")
GAISI_ZILS = PatternFill("solid", fgColor="2E75B6")
GALVINA = PatternFill("solid", fgColor="D9E1F2")
KOPSUMA = PatternFill("solid", fgColor="FFF2CC")
BALTS_BOLD = Font(color="FFFFFF", bold=True, size=12)
BOLD = Font(bold=True)
PLANA = Side(style="thin", color="999999")
RAMIS = Border(left=PLANA, right=PLANA, top=PLANA, bottom=PLANA)

VIENIBA_LBL = {"m3": "m³", "m2": "m²", "gab": "gab"}


def _virsraksts(ws, teksts, kolonnas, rinda=1, fill=ZILS):
    ws.merge_cells(start_row=rinda, start_column=1, end_row=rinda, end_column=kolonnas)
    c = ws.cell(rinda, 1, teksts)
    c.fill, c.font = fill, BALTS_BOLD
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[rinda].height = 22


def _galvina(ws, rinda, nosaukumi):
    for i, n in enumerate(nosaukumi, start=1):
        c = ws.cell(rinda, i, n)
        c.fill, c.font, c.border = GALVINA, BOLD, RAMIS
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _platumi(ws, platumi):
    for i, w in enumerate(platumi, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _baiti(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Pavadzīme ───────────────────────────────────────────────────────────────

def pavadzime_xlsx(pardosana: dict, rindas: list[dict], iestatijumi: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pavadzīme"
    _platumi(ws, [5, 30, 9, 9, 9, 8, 10, 10, 8, 11, 12])

    virsr = ("ATVĒRTAIS RĒĶINS (izraksts, nav galīgs)" if pardosana.get("statuss") == "atvērts"
             else "RĒĶINS-PAVADZĪME")
    _virsraksts(ws, f"{virsr} Nr. {pardosana['numurs']}", 11)
    ws.cell(2, 1, "Datums:").font = BOLD
    ws.cell(2, 3, pardosana["datums"])

    ws.cell(4, 1, "Pārdevējs:").font = BOLD
    ws.cell(4, 3, iestatijumi.get("pardevejs_nosaukums"))
    ws.cell(5, 1, "Reģ. Nr.:")
    ws.cell(5, 3, iestatijumi.get("pardevejs_regnr"))
    ws.cell(6, 1, "Adrese:")
    ws.cell(6, 3, iestatijumi.get("pardevejs_adrese"))
    ws.cell(7, 1, "Banka / konts:")
    ws.cell(7, 3, " ".join(x for x in [iestatijumi.get("pardevejs_banka"), iestatijumi.get("pardevejs_konts")] if x))
    ws.cell(8, 1, "Izsniegšanas vieta:")
    ws.cell(8, 3, iestatijumi.get("veikala_adrese"))

    ws.cell(4, 7, "Pircējs:").font = BOLD
    ws.cell(4, 8, pardosana.get("klients") or "Privātpersona")
    ws.cell(5, 7, "Reģ. Nr.:")
    ws.cell(5, 8, pardosana.get("klienta_regnr"))
    ws.cell(6, 7, "Adrese:")
    ws.cell(6, 8, pardosana.get("klienta_adrese"))
    ws.cell(7, 7, "PVN reģ. Nr.:")
    ws.cell(7, 8, pardosana.get("klienta_pvn_nr"))
    ws.cell(8, 7, "Apmaksa:")
    ws.cell(8, 8, pardosana.get("apmaksas_veids"))
    if pardosana.get("apmaksas_termins"):
        ws.cell(9, 7, "Apmaksas termiņš:")
        ws.cell(9, 8, pardosana["apmaksas_termins"])

    g = 11
    _galvina(ws, g, ["Nr.", "Produkts", "Biezums, mm", "Platums, mm", "Garums, mm", "Gab.",
                     "m³", "m²", "Vienība", "Cena, €", "Summa, €"])
    r = g
    for i, rd in enumerate(rindas, start=1):
        r = g + i
        ws.cell(r, 1, i)
        nos = rd["produkts"] + (f" ({rd['pakas_nr']})" if rd.get("pakas_nr") else "")
        ws.cell(r, 2, nos)
        ws.cell(r, 3, rd["biezums"])
        ws.cell(r, 4, rd["platums"])
        ws.cell(r, 5, rd["garums"])
        ws.cell(r, 6, rd["gab"])
        if rd["biezums"]:
            if rd["platums"]:
                ws.cell(r, 7, f"=ROUND(C{r}*D{r}*E{r}/1000000000*F{r},6)")
            else:  # apaļš — C ir diametrs
                ws.cell(r, 7, f"=ROUND(PI()*(C{r}/2)^2*E{r}/1000000000*F{r},6)")
            ws.cell(r, 8, f"=ROUND(D{r}*E{r}*F{r}/1000000,3)")
        ws.cell(r, 9, VIENIBA_LBL[rd["vieniba"]])
        ws.cell(r, 10, rd["cena"])
        daudz = {"m3": f"G{r}", "m2": f"H{r}", "gab": f"F{r}"}[rd["vieniba"]]
        ws.cell(r, 11, f"=ROUND({daudz}*J{r},2)")
        for k in range(1, 12):
            ws.cell(r, k).border = RAMIS
        ws.cell(r, 7).number_format = "0.000"
        ws.cell(r, 8).number_format = "0.000"
        ws.cell(r, 10).number_format = "0.00"
        ws.cell(r, 11).number_format = "#,##0.00"

    pirma, pedeja = g + 1, r
    k = pedeja + 1
    ws.cell(k, 2, "KOPĀ").font = BOLD
    ws.cell(k, 6, f"=SUM(F{pirma}:F{pedeja})")
    ws.cell(k, 7, f"=SUM(G{pirma}:G{pedeja})").number_format = "0.000"
    ws.cell(k, 8, f"=SUM(H{pirma}:H{pedeja})").number_format = "0.000"
    ws.cell(k, 11, f"=SUM(K{pirma}:K{pedeja})").number_format = "#,##0.00"
    for c in range(1, 12):
        ws.cell(k, c).fill = KOPSUMA
        ws.cell(k, c).border = RAMIS

    pvn_proc = pardosana.get("pvn_likme", 0.21)
    ws.cell(k + 2, 9, "Summa bez PVN:")
    ws.cell(k + 2, 11, f"=K{k}").number_format = "#,##0.00"
    ws.cell(k + 3, 9, "PVN likme:")
    ws.cell(k + 3, 11, pvn_proc).number_format = "0%"
    ws.cell(k + 4, 9, "PVN:")
    ws.cell(k + 4, 11, f"=ROUND(K{k + 2}*K{k + 3},2)").number_format = "#,##0.00"
    ws.cell(k + 5, 9, "KOPĀ APMAKSAI:").font = BOLD
    c = ws.cell(k + 5, 11, f"=K{k + 2}+K{k + 4}")
    c.number_format, c.font, c.fill = "#,##0.00", BOLD, KOPSUMA

    if not pvn_proc:
        ws.cell(k + 6, 2, REVERSA_TEKSTS).font = BOLD
    ws.cell(k + 8, 2, f"Izsniedza: {pardosana.get('pardevejs') or '____________________'}")
    ws.cell(k + 8, 7, "Saņēma: ____________________")
    if pardosana.get("piezimes"):
        ws.cell(k + 10, 2, f"Piezīmes: {pardosana['piezimes']}")

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return _baiti(wb)


# ── Atlikumi ────────────────────────────────────────────────────────────────

def atlikumi_xlsx(pakas: list[dict], datums: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Atlikumi"
    kol = ["Produkts", "Biezums", "Platums", "Garums", "Pakas Nr.", "Piegādātājs", "Atlikums, gab",
           "m³", "m²", "Pašizmaksa €/m³", "Pašizmaksa €/gab", "Pašizmaksas summa",
           "Cena €/m³", "Cena €/m²", "Cena €/gab", "Pārdošanas vērtība", "Faktiski saskaitīts", "Starpība"]
    _platumi(ws, [30, 8, 8, 8, 13, 15, 10, 9, 9, 11, 11, 13, 10, 10, 10, 13, 11, 10])
    _virsraksts(ws, f"JELGAVAS VEIKALS — NOLIKTAVAS ATLIKUMI uz {datums}", len(kol))
    _galvina(ws, 3, kol)
    r = 3
    for i, p in enumerate(pakas, start=4):
        r = i
        vals = [p["produkts"], p["biezums"], p["platums"], p["garums"], p["pakas_nr"], p["piegadatajs"],
                p["gab_atlikums"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, v)
        ws.cell(r, 8, f'=IF(B{r}="",0,IF(C{r}="",ROUND(PI()*(B{r}/2)^2*D{r}/1000000000*G{r},3),'
                               f'ROUND(B{r}*C{r}*D{r}/1000000000*G{r},3)))').number_format = "0.000"
        ws.cell(r, 9, f'=IF(C{r}="",0,ROUND(C{r}*D{r}*G{r}/1000000,3))').number_format = "0.000"
        ws.cell(r, 10, p["pasizmaksa_m3"])
        ws.cell(r, 11, p["pasizmaksa_gab"])
        ws.cell(r, 12, f"=ROUND(H{r}*N(J{r})+G{r}*N(K{r}),2)").number_format = "#,##0.00"
        ws.cell(r, 13, p["cena_m3"])
        ws.cell(r, 14, p["cena_m2"])
        ws.cell(r, 15, p["cena_gab"])
        # Vērtība: m² cena, ja ir; citādi m³ cena; citādi gab cena
        ws.cell(r, 16, f"=ROUND(IF(N(N{r})>0,I{r}*N{r},IF(N(M{r})>0,H{r}*M{r},G{r}*N(O{r}))),2)"
                ).number_format = "#,##0.00"
        ws.cell(r, 18, f'=IF(Q{r}="","",Q{r}-G{r})')
        for j in range(1, len(kol) + 1):
            ws.cell(r, j).border = RAMIS
    pirma, pedeja = 4, r
    k = pedeja + 1
    ws.cell(k, 1, "KOPĀ").font = BOLD
    for kolonna in ("G", "H", "I", "L", "P"):
        ws[f"{kolonna}{k}"] = f"=SUM({kolonna}{pirma}:{kolonna}{pedeja})"
    for j in range(1, len(kol) + 1):
        ws.cell(k, j).fill = KOPSUMA
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(kol))}{pedeja}"
    return _baiti(wb)


# ── Atskaite ────────────────────────────────────────────────────────────────

def atskaite_xlsx(rindas: list[dict], no: str, lidz: str, pardosanas: list[dict] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pārdošanas"
    kol = ["Datums", "Nr.", "Klients", "Apmaksa", "Produkts", "Biezums", "Platums", "Garums",
           "Pakas Nr.", "Piegādātājs", "Gab.", "m³", "Vienība", "Cena", "Summa bez PVN", "Pašizmaksa", "Delta",
           "Vēlamā cena", "Vēlamā summa", "Atlaide"]
    _platumi(ws, [11, 13, 20, 11, 28, 8, 8, 8, 12, 15, 7, 9, 8, 9, 12, 12, 11, 10, 12, 10])
    _virsraksts(ws, f"JELGAVAS VEIKALS — PĀRDOŠANAS {no} – {lidz}", len(kol))
    _galvina(ws, 3, kol)
    r = 3
    for i, rd in enumerate(rindas, start=4):
        r = i
        vals = [rd["datums"], rd["numurs"], rd["klients"], rd["apmaksas_veids"], rd["produkts"],
                rd["biezums"], rd["platums"], rd["garums"], rd["pakas_nr"], rd["piegadatajs"] or "—", rd["gab"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, v)
        ws.cell(r, 12, f'=IF(F{r}="",0,IF(G{r}="",ROUND(PI()*(F{r}/2)^2*H{r}/1000000000*K{r},6),'
                               f'ROUND(F{r}*G{r}*H{r}/1000000000*K{r},6)))').number_format = "0.000"
        ws.cell(r, 13, VIENIBA_LBL[rd["vieniba"]])
        ws.cell(r, 14, rd["cena"])
        ws.cell(r, 15, rd["summa"]).number_format = "#,##0.00"
        ws.cell(r, 16, rd["pasizmaksa"]).number_format = "#,##0.00"
        ws.cell(r, 17, f"=O{r}-P{r}").number_format = "#,##0.00"
        ws.cell(r, 18, rd.get("cena_velama") or rd["cena"])
        ws.cell(r, 19, rd.get("summa_velama") or rd["summa"]).number_format = "#,##0.00"
        ws.cell(r, 20, f"=S{r}-O{r}").number_format = "#,##0.00"
    pirma, pedeja = 4, max(r, 4)
    k = pedeja + 1
    ws.cell(k, 1, "KOPĀ").font = BOLD
    for kolonna in ("K", "L", "O", "P", "Q", "S", "T"):
        ws[f"{kolonna}{k}"] = f"=SUM({kolonna}{pirma}:{kolonna}{pedeja})"
    for j in range(1, len(kol) + 1):
        ws.cell(k, j).fill = KOPSUMA
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(kol))}{pedeja}"

    # Kopsavilkums pa piegādātājiem — SUMIF formulas no pirmās lapas
    ks = wb.create_sheet("Kopsavilkums")
    _platumi(ks, [22, 12, 12, 15, 15, 13])
    _virsraksts(ks, f"KOPSAVILKUMS PA PIEGĀDĀTĀJIEM {no} – {lidz}", 6)
    _galvina(ks, 3, ["Piegādātājs", "Pārdotie gab", "Pārdotie m³", "Summa bez PVN", "Pašizmaksa", "Delta"])
    piegadataji = sorted({rd["piegadatajs"] or "—" for rd in rindas})
    rng = lambda c: f"'Pārdošanas'!${c}${pirma}:${c}${pedeja}"
    rr = 3
    for rr, pg in enumerate(piegadataji, start=4):
        ks.cell(rr, 1, pg)
        for j, c in enumerate(("K", "L", "O", "P"), start=2):
            ks.cell(rr, j, f"=SUMIF({rng('J')},A{rr},{rng(c)})")
        ks.cell(rr, 6, f"=D{rr}-E{rr}")
        for j in range(1, 7):
            ks.cell(rr, j).border = RAMIS
    kk = rr + 1
    ks.cell(kk, 1, "KOPĀ").font = BOLD
    for j, c in enumerate("BCDEF", start=2):
        ks.cell(kk, j, f"=SUM({c}4:{c}{max(rr, 4)})")
    for j in range(1, 7):
        ks.cell(kk, j).fill = KOPSUMA
    for row in ks.iter_rows(min_row=4, max_row=kk, min_col=3, max_col=6):
        for c in row:
            c.number_format = "#,##0.000" if c.column == 3 else "#,##0.00"

    # Pa apmaksas veidiem (kases slēgšanai) — no čeku saraksta
    if pardosanas:
        ka = wb.create_sheet("Pa apmaksas veidiem")
        _platumi(ka, [14, 22, 20, 12, 10, 13, 12, 13])
        _virsraksts(ka, f"ČEKI UN APMAKSA {no} – {lidz}", 8)
        _galvina(ka, 3, ["Nr.", "Pircējs", "Apmaksa", "Datums", "PVN likme", "Bez PVN", "PVN", "Ar PVN"])
        r = 3
        for r, p in enumerate(pardosanas, start=4):
            for j, v in enumerate([p["numurs"], p["klients"] or "Privātpersona", p["apmaksas_veids"] or "—",
                                   p["datums"], p["pvn_likme"], p["summa_bez_pvn"]], start=1):
                ka.cell(r, j, v)
            ka.cell(r, 5).number_format = "0%"
            ka.cell(r, 7, f"=ROUND(F{r}*E{r},2)")
            ka.cell(r, 8, f"=F{r}+G{r}")
        pedeja = max(r, 4)
        kk = pedeja + 2
        _galvina(ka, kk, ["Apmaksa", "Čeki", "", "", "", "Bez PVN", "PVN", "Ar PVN"])
        veidi = sorted({p["apmaksas_veids"] or "—" for p in pardosanas})
        for i, v in enumerate(veidi, start=kk + 1):
            ka.cell(i, 1, v)
            ka.cell(i, 2, f"=COUNTIF($C$4:$C${pedeja},A{i})")
            for kol in "FGH":
                ka[f"{kol}{i}"] = f"=SUMIF($C$4:$C${pedeja},$A{i},{kol}$4:{kol}${pedeja})"
        kopa = kk + len(veidi) + 1
        ka.cell(kopa, 1, "KOPĀ").font = BOLD
        ka.cell(kopa, 2, f"=SUM(B{kk + 1}:B{kopa - 1})")
        for kol in "FGH":
            ka[f"{kol}{kopa}"] = f"=SUM({kol}{kk + 1}:{kol}{kopa - 1})"
        for j in range(1, 9):
            ka.cell(kopa, j).fill = KOPSUMA
        for row in ka.iter_rows(min_row=4, max_row=kopa, min_col=6, max_col=8):
            for c in row:
                c.number_format = "#,##0.00"
    return _baiti(wb)
